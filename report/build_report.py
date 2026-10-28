import subprocess, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = '/home/user/webapp'
FIG = f'{ROOT}/report/figures'
OUT = f'{ROOT}/report/ADMIT-Lipa_Prototype_Progress_Report.docx'
doc = Document()
for s in doc.sections:
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.top_margin = s.bottom_margin = Cm(2.54)
    s.left_margin = s.right_margin = Cm(2.54)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for n in ('Heading 1', 'Heading 2', 'Heading 3'):
    h = doc.styles[n]
    h.font.name = 'Times New Roman'
    h.font.color.rgb = RGBColor(0, 0, 0)
    h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 2'].font.size = Pt(12)

def center(text, bold=False, size=12, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p

def para(text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    return p

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hex_)
    tcPr.append(sh)

def table(headers, rows, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        shade(c, 'D9D9D9')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
    if widths:
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
            gc.set(qn('w:w'), str(int(widths[i] * 567)))
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t

def figure(path, caption, width=15.5):
    doc.add_picture(path, width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption); r.italic = True; r.font.size = Pt(10)

commit = subprocess.run(['git', '-C', ROOT, 'rev-parse', '--short', 'HEAD'], capture_output=True, text=True).stdout.strip()
test = json.load(open(f'{ROOT}/report/test_summary.json'))
video = json.load(open(f'{ROOT}/report/video_summary.json'))

center('ADMIT-LIPA:', True, 14)
center('A WEB-BASED ADMISSIONS CUSTOMER RELATIONSHIP', True, 14)
center('MANAGEMENT SYSTEM FOR THE ADMISSIONS OFFICE', True, 14)
center('OF STI COLLEGE LIPA', True, 14, after=24)
center('STI College Lipa', True)
center('C. M. Recto Avenue, Barangay 6, Lipa City, Batangas')
center('School of Information and Communications Technology', after=24)
center('Prototype Progress Report', True, 13, after=24)
center('Presented in Partial Fulfillment of the Requirements in')
center('Application Development and Emerging Technologies', after=30)
center('Submitted by the Proponents:', True)
for n in ['Atienza, Kurt Judd P.', 'Garonia, Ian Humprey M.', 'Apura, Alexis H.', 'Tenedo, Francis Gerald G.']:
    center(n)
center('', after=18)
center('Submitted to:', True)
center('Ms. Gabrielle Hernandez')
center('Course Facilitator', after=24)
center('First Semester, School Year 2026-2027')
center('Report date: October 28, 2026')
doc.add_page_break()

doc.add_heading('1. Summary of Progress', 1)
para('The proponents moved ADMIT-Lipa from the approved System Concept Proposal to a working web application. '
     'All fourteen use cases in the use-case diagram (Figure 1 of the proposal) now run end to end in the development build, '
     f'and the automated acceptance suite passed {test["passed"]} of {test["total"]} checks on October 28, 2026. '
     'Every functional requirement FR-01 to FR-14 is exercised by at least one check that follows the verification method written in the proposal.')
table(['Item', 'Status as of October 28, 2026'], [
    ['Use cases implemented', '14 of 14 (UC1 to UC14)'],
    ['Functional requirements verified by automated test', '14 of 14 (FR-01 to FR-14)'],
    ['Non-functional requirements verified by automated test', 'NFR-02, NFR-03 (in part); NFR-05 by screenshot at 390 px'],
    ['Screens designed and captured', '13 screens, desktop (1920 x 1080) and mobile (390 x 844), light and dark'],
    ['Source code', 'https://github.com/Tiredicey/APPDEV'],
    ['Deployment', 'Runs in the development environment; public hosting on Cloudflare Pages is the next step'],
], [7, 9])

doc.add_heading('2. Technology Used', 1)
table(['Layer', 'Choice', 'Reason'], [
    ['Server', 'Hono web framework on Cloudflare Workers / Pages', 'Runs at the network edge; no server for the campus to maintain'],
    ['Database', 'Cloudflare D1 (SQLite)', 'Relational tables for tickets, requirements, payments and the audit log; Cloudflare states D1 encrypts data at rest and in transit (Cloudflare, 2026)'],
    ['Field encryption', 'AES-256-GCM via the Web Crypto API', 'Date of birth, mobile number, address and financial-aid answers are encrypted before they reach the database (NFR-03)'],
    ['Passwords', 'PBKDF2-SHA-256, 100,000 iterations, per-user salt', 'No plain-text passwords stored'],
    ['Front end', 'Plain JavaScript single-page app, IBM Plex type, no framework', 'Small download for mobile users on slow connections'],
    ['Development assistant', 'Genspark AI coding agent', 'Generated code, tests and the background video under the proponents\' direction; the team reviewed the screens and tested the prototype'],
], [3, 5, 8])

doc.add_heading('3. Requirement Traceability: What Was Built and How It Was Verified', 1)
rows = [
    ['FR-01', 'UC1', 'Inquiry form returns a ticket number such as ADM-2026-00001 in the same session and stores it', 'PASS'],
    ['FR-02', 'UC2', 'Second profile with the same surname, first name and date of birth is blocked with a message (keyed HMAC index, so the duplicate check works on encrypted data)', 'PASS'],
    ['FR-03', 'UC3', 'Each required document shows received, pending or rejected; a rejection requires a reason', 'PASS'],
    ['FR-04', 'UC4', 'TES, TES-3A, TDP and SHS Voucher flagged from the encoded profile; boundary tested at P400,000 and P400,001', 'PASS'],
    ['FR-05', 'UC5', 'Applicant sees own stage only; opening another record returns "forbidden"', 'PASS'],
    ['FR-06', 'UC6', 'Elapsed working days computed against the 3 / 7 / 20-day tiers, excluding weekends and holidays entered by the administrator', 'PASS'],
    ['FR-07', 'UC7', 'Volume, stage conversion and case aging for any date range; totals reconcile with the raw case list; CSV export', 'PASS'],
    ['FR-08', 'UC8', 'Consent statement text, timestamp and purpose stored for every inquiry and application; no consent, no record', 'PASS'],
    ['FR-09', 'UC9', 'Reservation fee validated only with an official receipt number; a receipt number cannot be reused', 'PASS'],
    ['FR-10', 'UC10', 'Deficiency notice lists each document still pending after its deadline; repeat runs do not duplicate notices', 'PASS'],
    ['FR-11', 'UC11', 'Endorsed case reaches the financial aid coordinator queue with the program document list', 'PASS'],
    ['FR-12', 'UC12', 'Opening a record requires a stated purpose; user, record, time and purpose are logged; applicants can download their own access log', 'PASS'],
    ['FR-13', 'UC13', 'Ticket owned by one named officer; cannot be resolved without a closing remark', 'PASS'],
    ['FR-14', 'UC14', 'Overdue ticket raised to the School Administrator queue; on-time tickets are not', 'PASS'],
]
table(['Req.', 'Use case', 'Implemented behavior', 'Result'], rows, [1.6, 1.7, 11, 1.7], size=9)

doc.add_heading('4. Screens', 1)
para('The screens below were captured from the running build with an automated browser. '
     'The landing page uses the 1921 x 1080 "Landing Page" frame and its #FFF5C3 fill from the team Figma file as its base.')
for f, cap in [
    ('01-landing.png', 'Figure 2. Landing page with the looping background video (desktop).'),
    ('05-applicant-dashboard.png', 'Figure 3. Applicant dashboard: stage, requirements, subsidy matches, questions, consent (UC3, UC4, UC5, UC8).'),
    ('06-tickets.png', 'Figure 4. Officer ticket queue with service-level clock (UC6, UC13).'),
    ('13-applicant-record.png', 'Figure 5. Staff view of an applicant record: compliance, payment, aid endorsement (UC3, UC9, UC10, UC11).'),
    ('09-analytics.png', 'Figure 6. Admissions analytics for the School Administrator (UC7).'),
    ('10-audit.png', 'Figure 7. Personal data access log for the Data Protection Officer (UC12).'),
]:
    figure(f'{FIG}/{f}', cap)

doc.add_heading('5. Background Video', 1)
para(f'The landing page plays a {video["duration"]:.0f}-second seamless loop at {video["fps"]} frames per second, {video["width"]} x {video["height"]}, '
     f'{video["size_mb"]:.1f} MB, H.264. It shows a slow fall down the core of an abandoned orbital station. '
     f'Camera speed and the amber warning lamps pulse at 140 beats per minute. The largest frame-to-frame brightness change is {video["max_delta"]:.2f} of 255 levels, far below the WCAG 2.3.1 flash threshold. '
     'The video is muted, carries no information, and has a Pause control. It does not autoplay for visitors who set "reduce motion" on their device.')

doc.add_heading('6. Member Contributions', 1)
para('Each member signs his own row to confirm the contribution listed.')
contrib = table(['Member', 'Role', 'Contribution', 'Signature'], [
    ['Atienza, Kurt Judd P.', 'Project Lead', 'Designed the landing frame in Figma; directed the AI-assisted build of the back end, database and test suite; integrated and tested the prototype', ''],
    ['Garonia, Ian Humprey M.', 'UX/UI Designer and Tester', 'Designed screens in Figma; tested the prototype', ''],
    ['Apura, Alexis H.', 'UX/UI Designer and Tester', 'Designed screens in Figma; tested the prototype', ''],
    ['Tenedo, Francis Gerald G.', 'UX/UI Designer and Tester', 'Designed screens in Figma; tested the prototype', ''],
], [4, 3.4, 6, 3], size=10)
for row in contrib.rows[1:]:
    row.height = Cm(1.4)

doc.add_heading('References', 1)
for ref in [
    'Cloudflare. (2026, April 21). Data security. Cloudflare D1 docs. https://developers.cloudflare.com/d1/reference/data-security/',
    'Department of Education. (n.d.). Senior high school voucher program guidelines. https://depedroxi.ph/voucher-program/',
    'STI Education Services Group. (n.d.). Admissions. https://www.sti.edu/admissions.asp',
    'Unified Student Financial Assistance System for Tertiary Education. (n.d.). Tertiary education subsidy and Tulong Dunong program. https://unifast.gov.ph/tes.html',
    'World Wide Web Consortium. (2023). Web Content Accessibility Guidelines (WCAG) 2.2. https://www.w3.org/TR/WCAG22/',
]:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)

doc.save(OUT)
print('ok')
